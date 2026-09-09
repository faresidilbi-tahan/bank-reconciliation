"""
POST /api/extract-bank
Body: { "file_base64": "<base64>", "filename": "statement.pdf" }

Deterministic, format-agnostic BANK statement parser - no LLM, no
external calls. Distinct from the supplier-statement parser this project
is modeled on in one structural way: banks commonly print a single
SIGNED "Amount" column (negative = money out, positive = money in)
instead of separate Debit/Credit columns. This module detects that case
and splits the signed amount into debit/credit itself, so every
downstream row has the same shape (date/id/description/debit/credit/
balance) regardless of which column style the source used.

Real quirks handled, found from an actual BLOM Bank statement:
  - A two-line wrapped column header ("Business" / "Date" stacked across
    two physical lines) sits right next to an unrelated "Value Date"
    column that ALSO contains the word "Date" - naive keyword matching
    would anchor to Value Date (a settlement date, not the transaction
    date) instead of the real transaction date. Detected by looking for
    the word "Business" near the header and treating its position as the
    true date anchor, overriding any other "date" match nearby.
  - A trailing "Pending Transactions" section (not-yet-cleared items)
    after all real transactions, in a completely different table layout -
    excluded entirely once that marker is seen, the same way a trailing
    postdated-cheques table is excluded on the supplier side.
"""

from http.server import BaseHTTPRequestHandler
import base64
import csv
import io
import json
import re

ARABIC_RUN_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F][\u0600-\u06FF\u0750-\u077F\s]*[\u0600-\u06FF\u0750-\u077F]|[\u0600-\u06FF\u0750-\u077F]")


def fix_bidi_text(text):
    return ARABIC_RUN_RE.sub(lambda m: m.group(0)[::-1], text)


import datetime as dt

import pdfplumber
import openpyxl

BUILD_TAG = "2026-08-27-bank-v1"

# "amount" is new here vs the supplier parser - a single signed column
# instead of separate debit/credit. "id" here means whatever reference
# number the bank prints (transaction ref, cheque number) - captured for
# DISPLAY only; per this project's matching rules, id is never used to
# match rows (amount + date only, see compare.py).
COLUMN_KEYWORDS = {
    "date": ["date", "invc", "تاريخ", "التاريخ"],
    "id": ["ref", "reference", "transaction ref", "trx", "cheque", "check",
           "chq", "voucher", "no", "number", "رقم", "مرجع"],
    "description": ["narrative", "description", "details", "particulars",
                    "memo", "remarks", "بيان", "البيان", "التفاصيل", "الوصف"],
    "amount": ["amount", "مبلغ", "المبلغ"],
    "debit": ["debit", "withdrawal", "dr", "مدين"],
    "credit": ["credit", "deposit", "cr", "دائن"],
    "balance": ["balance", "bal", "رصيد", "الرصيد"],
}

OPENING_WORDS = ("opening", "b/f", "b.f", "brf", "brought", "balance until",
                 "افتتاحي", "سابق", "مدور", "رصيد اول")
CLOSING_WORDS = ("closing", "c/f", "c.f", "carried", "ending balance",
                 "end date", "balance as at",
                 "اجمالي", "إجمالي", "المجموع", "ختامي", "نهائي")
SKIP_WORDS = ("statement", "page ", "page:", "printed", "tel:", "fax:",
              "p.o.box", "www.", "@")

# Once seen, everything after is trailing metadata, never a real
# transaction - a "Pending Transactions" section on a real BLOM Bank
# statement used a completely different table layout (Date/Description/
# Merchant Name/Amount instead of the main table's columns) and its rows
# aren't cleared/final, so they must never be reconciled as real activity.
PENDING_SECTION_RE = re.compile(r"pending\s*transactions?", re.IGNORECASE)

AMOUNT_RE = re.compile(r"^\(?-?(?:[\d,]+(?:\.\d+)?|\.\d+)\)?(CR|DR|DB)?$", re.IGNORECASE)
NUM_DATE_RE = re.compile(r"(\d{1,4})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,4})")
MONTH_DATE_RE = re.compile(
    r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+(\d{1,2})\s*,?\s+(\d{4})",
    re.IGNORECASE)
MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
          "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}


def match_column(text):
    t = str(text or "").strip().lower().strip(":.")
    for col, words in COLUMN_KEYWORDS.items():
        if t in words:
            return col
    parts = [p.strip(":.") for p in t.split()]
    for col, words in COLUMN_KEYWORDS.items():
        if any(p in words for p in parts):
            return col
    return None


def parse_amount(token):
    token = str(token).strip()
    cleaned = re.sub(r"(?i)us\$|usd|\$", "", token).strip()
    m = AMOUNT_RE.match(cleaned)
    if not m:
        return None
    s = cleaned[: len(cleaned) - 2] if m.group(1) else cleaned
    s = s.replace(",", "").replace("(", "-").replace(")", "")
    try:
        return float(s)
    except ValueError:
        return None


def parse_amount_cell(raw_cell):
    if raw_cell == "":
        return None
    val = parse_amount(raw_cell)
    if val is not None:
        return val
    parts = raw_cell.split()
    while parts and re.fullmatch(r"[A-Za-z]{2,3}", parts[-1]):
        parts = parts[:-1]
    if not parts:
        return None
    return parse_amount(parts[-1])


_DATE_CONVENTION = "DMY"


def detect_date_convention(text):
    dmy_evidence = mdy_evidence = 0
    for m in NUM_DATE_RE.finditer(text):
        a, b = int(m.group(1)), int(m.group(2))
        if a > 31 or b > 31:
            continue
        if a > 12 and b <= 12:
            dmy_evidence += 1
        elif b > 12 and a <= 12:
            mdy_evidence += 1
    return "MDY" if mdy_evidence > dmy_evidence else "DMY"


def find_date(text):
    text = str(text or "")
    m = MONTH_DATE_RE.search(text)
    if m:
        mo = MONTHS[m.group(1).lower()[:3]]
        d, y = int(m.group(2)), int(m.group(3))
        if 1 <= d <= 31:
            return f"{y:04d}-{mo:02d}-{d:02d}"
    m = NUM_DATE_RE.search(text)
    if m:
        a, b, c = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if a > 31:
            y, mo, d = a, b, c
        elif a > 12 and b <= 12:
            d, mo, y = a, b, c
        elif b > 12 and a <= 12:
            mo, d, y = a, b, c
        elif _DATE_CONVENTION == "MDY":
            mo, d, y = a, b, c
        else:
            d, mo, y = a, b, c
        if y < 100:
            y += 2000
        if 1 <= d <= 31 and 1 <= mo <= 12 and 1900 < y < 2100:
            return f"{y:04d}-{mo:02d}-{d:02d}"
    return None


def build_row(cells, raw_low):
    date = find_date(cells.get("date", "")) or ""

    debit = parse_amount_cell(cells.get("debit", ""))
    credit = parse_amount_cell(cells.get("credit", ""))
    if debit is None and credit is None:
        # Single signed "Amount" column - negative is money out (debit),
        # positive is money in (credit). This is the common bank-
        # statement shape this parser is specifically built to support.
        amount = parse_amount_cell(cells.get("amount", ""))
        if amount is not None:
            if amount < 0:
                debit, credit = abs(amount), 0.0
            else:
                debit, credit = 0.0, amount
    balance = parse_amount_cell(cells.get("balance", ""))

    if any(k in raw_low for k in OPENING_WORDS):
        row_type = "opening_balance"
    elif any(k in raw_low for k in CLOSING_WORDS):
        row_type = "closing_balance"
    else:
        row_type = "transaction"

    if not (date or any(v is not None for v in (debit, credit, balance))):
        return None

    desc = fix_bidi_text(cells.get("description", ""))
    return {
        "date": date,
        "id": str(cells.get("id", "") or "").strip(),
        "description": desc,
        "debit": abs(debit) if debit is not None else 0.0,
        "credit": abs(credit) if credit is not None else 0.0,
        "balance": balance if balance is not None else 0.0,
        "row_type": row_type,
    }


def is_continuation(cells):
    if find_date(cells.get("date", "")):
        return False
    if any(parse_amount(cells.get(c, "")) is not None for c in ("debit", "credit", "balance", "amount") if cells.get(c, "") != ""):
        return False
    return bool(cells.get("description") or cells.get("id"))


# ================================================================== PDF

def group_lines(page):
    words = page.extract_words(x_tolerance=1.5, y_tolerance=2.0, keep_blank_chars=False)
    words.sort(key=lambda w: (w["top"], w["x0"]))
    lines, cur, cur_top = [], [], None
    for w in words:
        if cur_top is None or abs(w["top"] - cur_top) <= 2.5:
            cur.append(w)
            if cur_top is None:
                cur_top = w["top"]
        else:
            lines.append(sorted(cur, key=lambda x: x["x0"]))
            cur, cur_top = [w], w["top"]
    if cur:
        lines.append(sorted(cur, key=lambda x: x["x0"]))
    return lines


def find_header(lines):
    """Same general keyword-anchor approach as the supplier parser, plus
    one bank-specific override: if a word literally reading "Business"
    appears near the header (the "Business Date"/"Value Date" pattern
    seen on a real statement, where BOTH column labels contain the word
    "Date"), its position wins as the date anchor - "Business Date" is
    the actual transaction date; "Value Date" is a secondary settlement
    date this project has no use for and would otherwise be picked
    instead, since a generic scan has no way to prefer one "Date" over
    another."""
    best, best_count, best_idx = None, 0, None
    for idx, line in enumerate(lines):
        cols = {}
        for w in line:
            col = match_column(w["text"])
            if col and col not in cols:
                cols[col] = (w["x0"] + w["x1"]) / 2.0
        if len(cols) >= 3 and any(c in cols for c in ("debit", "credit", "balance", "amount")):
            if len(cols) > best_count:
                best, best_count, best_idx = cols, len(cols), idx
    if best is None:
        return None

    for offset in (0, -1, 1):
        check_idx = best_idx + offset
        if 0 <= check_idx < len(lines):
            for w in lines[check_idx]:
                if w["text"].strip().lower() == "business":
                    best["date"] = (w["x0"] + w["x1"]) / 2.0
                    return best
    return best


def build_intervals(anchors, page_width):
    ordered = sorted(anchors.items(), key=lambda kv: kv[1])
    intervals = []
    for idx, (col, x) in enumerate(ordered):
        left = 0 if idx == 0 else (ordered[idx - 1][1] + x) / 2.0
        right = page_width if idx == len(ordered) - 1 else (x + ordered[idx + 1][1]) / 2.0
        intervals.append((col, left, right))
    return intervals


def assign_columns(line, intervals):
    cells = {col: [] for col, _, _ in intervals}
    for w in line:
        center = (w["x0"] + w["x1"]) / 2.0
        for col, left, right in intervals:
            if left <= center < right:
                cells[col].append(w["text"].strip())
                break
    return {col: " ".join(v).strip() for col, v in cells.items()}


def word_column(w, intervals):
    center = (w["x0"] + w["x1"]) / 2.0
    for col, left, right in intervals:
        if left <= center < right:
            return col
    return None


def stitch_wrapped_negative_amounts(lines, intervals):
    """A negative amount too wide for its line sometimes renders as just
    its "-" sign in place, with the magnitude pushed onto the START of
    the NEXT physical line while balance/ref/continuation text on the
    original line renders normally around it (seen for real: "Trf To
    009/02/353/0695777/1 - 4,259.48 1260407 : . . . . . . Beneficiary"
    then, starting the next line, "45,000.00 009-02-353-0695777-1-8" -
    45,000.00 is the missing magnitude, landing exactly in the amount
    column's x-range on the next line down). Detected by column position,
    not line order: a lone "-" sitting alone in the amount column on one
    line, paired with a parseable number sitting in that SAME column
    range on the next line."""
    amount_range = next(((l, r) for c, l, r in intervals if c == "amount"), None)
    if not amount_range:
        return lines
    lo, hi = amount_range
    out = []
    i = 0
    while i < len(lines):
        line = lines[i]
        minus_words = [w for w in line if w["text"].strip() == "-" and lo <= (w["x0"] + w["x1"]) / 2 < hi]
        other_in_col = [w for w in line if w not in minus_words and lo <= (w["x0"] + w["x1"]) / 2 < hi]
        if minus_words and not other_in_col and i + 1 < len(lines):
            nxt = lines[i + 1]
            candidate = next(
                (w for w in nxt if lo <= (w["x0"] + w["x1"]) / 2 < hi and parse_amount(w["text"]) is not None),
                None)
            if candidate:
                minus_word = minus_words[0]
                merged = dict(candidate)
                merged["text"] = "-" + candidate["text"]
                merged["x0"], merged["x1"] = minus_word["x0"], minus_word["x1"]
                new_line = sorted([w for w in line if w is not minus_word] + [merged], key=lambda w: w["x0"])
                remainder = [w for w in nxt if w is not candidate]
                out.append(new_line)
                if remainder:
                    lines[i + 1] = remainder
                else:
                    i += 1
                i += 1
                continue
        out.append(line)
        i += 1
    return out


def parse_words_strategy(pdf):
    rows, warnings = [], []
    anchors = None
    seen_pending = False
    for page_no, page in enumerate(pdf.pages, start=1):
        if seen_pending:
            break
        lines = group_lines(page)
        # Never let a DIFFERENT table's header (like the trailing Pending
        # Transactions section's own "Date Description Merchant Name
        # Amount" line, seen for real on this document's last page)
        # compete for this page's column anchors - it uses different
        # column positions entirely, and would silently corrupt every
        # real transaction line above it on the same page if it won.
        header_scan_lines = lines
        for li, line in enumerate(lines):
            if PENDING_SECTION_RE.search(" ".join(w["text"] for w in line)):
                header_scan_lines = lines[:li]
                break
        page_anchors = find_header(header_scan_lines)
        if page_anchors:
            anchors = page_anchors
        if not anchors:
            continue
        intervals = build_intervals(anchors, page.width)
        lines = stitch_wrapped_negative_amounts(lines, intervals)

        prev_was_data = False
        prev_date = ""
        for line in lines:
            if seen_pending:
                break
            raw = " ".join(w["text"] for w in line).strip()
            if not raw:
                continue
            raw_low = raw.lower()

            if PENDING_SECTION_RE.search(raw_low):
                seen_pending = True
                break

            # A sentence-style opening/closing balance line ("Brought
            # Forward Balance: USD 11,356.05 C") isn't real tabular data -
            # its wording doesn't align with the transaction table's own
            # column positions, so slicing it by those intervals scrambles
            # it across the wrong cells. Pull the balance value straight
            # from the raw line text instead of going through column
            # assignment at all.
            if any(k in raw_low for k in OPENING_WORDS) or any(k in raw_low for k in CLOSING_WORDS):
                amounts_found = [a for a in (parse_amount(w["text"]) for w in line) if a is not None]
                rows.append({
                    "date": "", "id": "", "description": raw,
                    "debit": 0.0, "credit": 0.0,
                    "balance": amounts_found[-1] if amounts_found else 0.0,
                    "row_type": "opening_balance" if any(k in raw_low for k in OPENING_WORDS) else "closing_balance",
                })
                prev_was_data = True
                continue

            header_hits = sum(
                1 for w in line
                if match_column(w["text"]) and match_column(w["text"]) == word_column(w, intervals)
            )
            if header_hits >= 3:
                prev_was_data = False
                continue
            if any(s in raw_low for s in SKIP_WORDS):
                prev_was_data = False
                continue

            cells = assign_columns(line, intervals)
            row = build_row(cells, raw_low)
            if row:
                if not row["date"] and prev_date and row["row_type"] == "transaction":
                    row["date"] = prev_date
                if row["row_type"] == "transaction" and not row["date"]:
                    prev_was_data = False
                    continue
                if row["date"]:
                    prev_date = row["date"]
                rows.append(row)
                prev_was_data = True
            elif prev_was_data and is_continuation(cells) and rows:
                extra = (cells.get("description", "") + " " + cells.get("id", "")).strip()
                rows[-1]["description"] = (rows[-1]["description"] + " " + extra).strip()
            elif len(raw) > 3:
                warnings.append(f"page {page_no}: unclassified line: {raw[:120]}")
                prev_was_data = False
    return rows, warnings, anchors is not None


def parse_pdf(file_bytes):
    global _DATE_CONVENTION
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        pages = len(pdf.pages)
        total_chars = sum(len(p.chars) for p in pdf.pages)
        if total_chars < 20:
            raise ValueError(
                "This PDF has no text layer (it is a scan/image). "
                "It needs OCR before it can be parsed without an LLM."
            )
        full_text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        _DATE_CONVENTION = detect_date_convention(full_text)

        rows, warnings, found = parse_words_strategy(pdf)

    if not found:
        raise ValueError(
            "Could not find a recognizable column header (Date/Amount or "
            "Debit/Credit/Balance) in this PDF. This bank's format needs a "
            "keyword added to COLUMN_KEYWORDS."
        )
    warnings = [f"[pdf] {w}" for w in warnings]
    if not rows:
        warnings.append("[pdf] Header found but no data rows extracted.")
    return rows, warnings, {"pages": pages}


# ================================================================== XLSX

def cell_to_text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def map_sheet_header(header_row):
    mapped = {}
    for idx, cell in enumerate(header_row):
        col = match_column(cell)
        if col and col not in mapped:
            mapped[col] = idx
    return mapped


def build_sheet_row(row_values, mapped):
    def get(field):
        i = mapped.get(field, -1)
        return cell_to_text(row_values[i]) if 0 <= i < len(row_values) else ""

    cells = {f: get(f) for f in ("date", "id", "description", "amount", "debit", "credit", "balance")}
    raw_low = " ".join(cell_to_text(v) for v in row_values if v is not None).lower()
    return build_row(cells, raw_low), cells


def parse_xlsx(file_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    rows, warnings = [], []
    sheets_used = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue

        header_idx, mapped = None, None
        for i, r in enumerate(all_rows[:5]):
            m = map_sheet_header(r)
            if len(m) >= 3 and any(c in m for c in ("debit", "credit", "balance", "amount")):
                header_idx, mapped = i, m
                break
        if header_idx is None:
            continue

        sheets_used += 1
        prev_row_ref = None
        seen_pending = False
        for r in all_rows[header_idx + 1:]:
            if seen_pending:
                break
            if r is None or all(v is None for v in r):
                continue
            row_text = " ".join(cell_to_text(v) for v in r if v is not None).lower()
            if PENDING_SECTION_RE.search(row_text):
                seen_pending = True
                break
            row, cells = build_sheet_row(r, mapped)
            if row:
                rows.append(row)
                prev_row_ref = rows[-1]
            elif prev_row_ref is not None and is_continuation(cells):
                extra = (cells.get("description", "") + " " + cells.get("id", "")).strip()
                if extra:
                    prev_row_ref["description"] = (prev_row_ref["description"] + " " + extra).strip()

    if sheets_used == 0:
        raise ValueError(
            "Could not find a recognizable column header (Date/Amount or "
            "Debit/Credit/Balance) in any sheet of this workbook. This "
            "bank's format needs a keyword added to COLUMN_KEYWORDS."
        )
    if not rows:
        warnings.append("[xlsx] Header(s) found but no data rows extracted.")
    return rows, warnings, {"sheets": len(wb.sheetnames), "sheets_used": sheets_used}


# ================================================================== CSV

def parse_csv(file_bytes):
    global _DATE_CONVENTION
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Could not decode this CSV file as text.")
    _DATE_CONVENTION = detect_date_convention(text)

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = list(csv.reader(io.StringIO(text), dialect))
    reader = [r for r in reader if any(c.strip() for c in r)]
    if not reader:
        raise ValueError("This CSV file has no rows.")

    header_idx, mapped = None, None
    for i, r in enumerate(reader[:5]):
        m = map_sheet_header(r)
        if len(m) >= 3 and any(c in m for c in ("debit", "credit", "balance", "amount")):
            header_idx, mapped = i, m
            break
    if header_idx is None:
        raise ValueError(
            "Could not find a recognizable column header (Date/Amount or "
            "Debit/Credit/Balance) in this CSV. This bank's format needs "
            "a keyword added to COLUMN_KEYWORDS."
        )

    rows, warnings = [], []
    prev_row_ref = None
    seen_pending = False
    for r in reader[header_idx + 1:]:
        if seen_pending:
            break
        row_text = " ".join(str(c) for c in r).lower()
        if PENDING_SECTION_RE.search(row_text):
            seen_pending = True
            break
        row, cells = build_sheet_row(r, mapped)
        if row:
            rows.append(row)
            prev_row_ref = rows[-1]
        elif prev_row_ref is not None and is_continuation(cells):
            extra = (cells.get("description", "") + " " + cells.get("id", "")).strip()
            if extra:
                prev_row_ref["description"] = (prev_row_ref["description"] + " " + extra).strip()

    if not rows:
        warnings.append("[csv] Header found but no data rows extracted.")
    return rows, warnings, {}


# ================================================================== dispatch

def sniff_format(file_bytes, filename):
    if file_bytes.startswith(b"%PDF"):
        return "pdf"
    if file_bytes.startswith(b"PK\x03\x04"):
        return "xlsx"
    if file_bytes.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValueError(
            "This looks like a legacy .xls file (pre-2007 Excel format). "
            "Please re-save it as .xlsx or .csv and upload again."
        )
    ext = (filename or "").rsplit(".", 1)[-1].lower() if filename else ""
    if ext in ("csv", "tsv", "txt"):
        return "csv"
    if ext in ("xlsx", "xlsm"):
        return "xlsx"
    if ext == "pdf":
        return "pdf"
    try:
        file_bytes[:2048].decode("utf-8")
        return "csv"
    except UnicodeDecodeError:
        raise ValueError(
            "Could not identify this file's format. Supported formats: "
            "PDF, Excel (.xlsx), and CSV."
        )


def parse_bank_file(file_bytes, filename=None):
    fmt = sniff_format(file_bytes, filename)
    if fmt == "pdf":
        rows, warnings, meta = parse_pdf(file_bytes)
    elif fmt == "xlsx":
        rows, warnings, meta = parse_xlsx(file_bytes)
    else:
        rows, warnings, meta = parse_csv(file_bytes)
    result = {"rows": rows, "warnings": warnings, "build_tag": BUILD_TAG, "format": fmt}
    result.update(meta)
    return result


class handler(BaseHTTPRequestHandler):
    def _send(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        try:
            length = int(self.headers.get("Content-Length", 0))
            data = json.loads(self.rfile.read(length) or b"{}")
            b64 = data.get("file_base64") or data.get("pdf_base64", "")
            filename = data.get("filename", "")
            if "," in b64[:80]:
                b64 = b64.split(",", 1)[1]
            if not b64:
                return self._send(400, {"error": "file_base64 is required."})
            file_bytes = base64.b64decode(b64)
            self._send(200, parse_bank_file(file_bytes, filename))
        except json.JSONDecodeError:
            self._send(400, {"error": "Invalid JSON body."})
        except ValueError as e:
            self._send(422, {"error": str(e)})
        except Exception as e:  # noqa: BLE001
            self._send(500, {"error": f"Extraction failed: {e}"})
